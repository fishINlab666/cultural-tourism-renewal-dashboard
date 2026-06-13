from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "assets" / "项目基础数据模板.xlsx"

FIELDS = [
    ("项目名称", "榕城古厝活化示范项目", "必填", "填写项目正式名称"),
    ("项目类型", "历史街区活化", "必填", "从下拉选项中选择"),
    ("项目位置", "福建省福州市历史文化街区", "必填", "填写省、市、区及项目位置"),
    ("建筑/项目面积（㎡）", 4200, "必填", "填写 50-500000 之间的数值"),
    ("当前空置率（%）", 62, "必填", "填写 0-100 之间的数值"),
    ("保护与现状等级", "历史建筑", "必填", "从下拉选项中选择"),
    ("预算上限（万元）", 860, "必填", "填写 50-100000 之间的数值"),
    ("目标回本周期（年）", 3, "必填", "填写 0.5-20 之间的数值"),
    ("主要服务客群", "文旅主管部门", "必填", "从下拉选项中选择"),
    ("初步风格方向", "低干预修缮、在地文化、当代简约", "必填", "用关键词描述期望风格"),
    ("项目单位", "某文旅发展集团", "选填", "填写项目建设或运营单位"),
    ("现状与核心诉求", "盘活闲置院落，引入文化体验、轻餐饮和研学活动，同时控制文保、消防与夜间运营风险。", "选填", "简要说明现状和希望解决的问题"),
]


def style_header(row):
    for cell in row:
        cell.fill = PatternFill("solid", fgColor="1F4D3A")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    guide = workbook.active
    guide.title = "填写说明"
    guide.append(["项目数据模板填写说明"])
    guide["A1"].font = Font(size=18, bold=True, color="1F4D3A")
    guide.append(["1. 请只修改“项目基础数据”工作表中的“填写内容”列。"])
    guide.append(["2. 标记为必填的字段需要完整填写，数值字段不要填写单位。"])
    guide.append(["3. 上传后系统仅要求校对缺失、格式错误或超出范围的字段。"])
    guide.append(["4. 客户数据默认只用于当前项目，不自动进入集团案例库，也不用于模型训练。"])
    guide.column_dimensions["A"].width = 96

    sheet = workbook.create_sheet("项目基础数据")
    sheet.append(["字段", "填写内容", "是否必填", "填写说明"])
    style_header(sheet[1])
    for row in FIELDS:
        sheet.append(row)

    widths = {"A": 28, "B": 58, "C": 14, "D": 48}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:D{sheet.max_row}"
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    validations = {
        3: '"历史街区活化,古建院落活化,旧厂房文旅更新,景区配套提升,公共文化空间再运营"',
        7: '"一般存量建筑,历史建筑,文保建筑,历史文化街区"',
        10: '"文旅主管部门,国资/城投平台,景区运营方,文旅开发与商业运营方"',
    }
    for row_number, formula in validations.items():
        validation = DataValidation(type="list", formula1=formula, allow_blank=False)
        validation.error = "请选择下拉列表中的标准选项"
        validation.errorTitle = "填写内容不符合模板"
        sheet.add_data_validation(validation)
        validation.add(sheet[f"B{row_number}"])

    workbook.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
